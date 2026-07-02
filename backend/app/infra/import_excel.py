"""Importa o catalogo (modulos + inversores) da aba "BD" da planilha
Fotus oficial (`2026 - Dimensionamento Módulo x Inversor`) para o banco
SQLite local.

A aba BD tem duas Tabelas do Excel nomeadas (`mod` e `inv`) com
cabecalhos oficiais — lemos por NOME de coluna via essas Tabelas (nao
por letra de coluna hardcoded), o que torna a importacao resiliente a
insercao/remocao de colunas na planilha.

Marca (brand) e Modelo (model) ficam nas colunas B/C, FORA do range
das Tabelas nomeadas (que comecam em D) — lidas diretamente por
posicao fixa relativa a cada linha.

Uso:
    python -m app.infra.import_excel "caminho\\para\\planilha.xlsx"
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from app.infra.db import Base, SessionLocal, engine
from app.infra.models import InverterRow, ModuleRow

BRAND_COL = 2  # B
MODEL_COL = 3  # C

MPPT_PAIR_COUNT = 14


def _read_table_rows(ws: Worksheet, ws_values: Worksheet, table_ref: str) -> List[Dict[str, Any]]:
    min_col, min_row, max_col, max_row = range_boundaries(table_ref)
    header_row = min_row
    col_names = [ws.cell(row=header_row, column=c).value for c in range(min_col, max_col + 1)]
    rows: List[Dict[str, Any]] = []
    for r in range(min_row + 1, max_row + 1):
        row: Dict[str, Any] = {"_row": r}
        for idx, c in enumerate(range(min_col, max_col + 1)):
            row[col_names[idx]] = ws_values.cell(row=r, column=c).value
        row["brand"] = ws_values.cell(row=r, column=BRAND_COL).value
        row["model"] = ws_values.cell(row=r, column=MODEL_COL).value
        rows.append(row)
    return rows


def _to_float(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_modules(ws: Worksheet, ws_values: Worksheet) -> List[dict]:
    table_ref = ws.tables["mod"].ref
    rows = _read_table_rows(ws, ws_values, table_ref)
    modules: List[dict] = []
    for row in rows:
        pnom = _to_float(row.get("Pnom [Wp]"))
        if pnom is None or not row.get("brand"):
            continue  # linha separadora ("*** MARCA ***") ou vazia
        modules.append(
            {
                "brand": str(row["brand"]).strip(),
                "model": str(row["model"] or "").strip(),
                "label": f"{row['brand']} | {row['model']}",
                "pnom": pnom,
                "isc": _to_float(row.get("Isc [A]")) or 0.0,
                "voc": _to_float(row.get("Voc [V]")) or 0.0,
                "imp": _to_float(row.get("Imp [A]")) or 0.0,
                "vmp": _to_float(row.get("Vmp [V]")) or 0.0,
                "coef_v": _to_float(row.get("Coef V [%/ºC]")) or 0.0,
                "coef_i": _to_float(row.get("Coef I [%/°C]")) or 0.0,
                "coef_pmp": _to_float(row.get("Coef Pmp [%/ºC]")) or 0.0,
                "efic": _to_float(row.get("Eficiência (%)")),
                "altura_mm": _to_float(row.get("Altura (mm)")) or 0.0,
                "largura_mm": _to_float(row.get("Largura (mm)")) or 0.0,
                "espessura_mm": _to_float(row.get("Espessura (mm)")) or 0.0,
                "peso_kg": _to_float(row.get("Peso (kg)")) or 0.0,
            }
        )
    return modules


def parse_inverters(ws: Worksheet, ws_values: Worksheet) -> List[dict]:
    table_ref = ws.tables["inv"].ref
    rows = _read_table_rows(ws, ws_values, table_ref)
    inverters: List[dict] = []
    for row in rows:
        p_nom = _to_float(row.get("P nom c.a. [W]"))
        if p_nom is None or not row.get("brand"):
            continue

        mppt_currents = []
        for i in range(1, MPPT_PAIR_COUNT + 1):
            imax_key = "I max [A]" if i == 1 else f"I max ({i}) [A]"
            isc_key = "Isc max [A]" if i == 1 else f"Isc max ({i}) [A]"
            imax = _to_float(row.get(imax_key))
            isc = _to_float(row.get(isc_key))
            if imax is None and isc is None:
                continue
            mppt_currents.append({"imax": imax or 0.0, "isc": isc or 0.0})

        lin_max = []
        for i in range(1, MPPT_PAIR_COUNT + 1):
            key = "LIN MAX  1" if i == 1 else f"LIN MAX {i}"
            lin_max.append(_to_float(row.get(key)) or 0.0)

        num_mppt_raw = row.get("MPPT")
        num_mppt = int(num_mppt_raw) if num_mppt_raw not in (None, "") else len(mppt_currents)

        num_inputs_raw = row.get("ENTRADAS")
        num_inputs = int(num_inputs_raw) if num_inputs_raw not in (None, "") else None

        tensao = _to_float(row.get("Tensão"))
        fase = row.get("Fase") or None
        protection = row.get("Proteção CC") or None

        inverters.append(
            {
                "brand": str(row["brand"]).strip(),
                "model": str(row["model"] or "").strip(),
                "categoria": None,
                "p_nom": p_nom,
                "p_max_cc": _to_float(row.get("P max_cc [W]")),
                "v_max": _to_float(row.get("V max [V]")),
                "v_mpp_min": _to_float(row.get("V MPPT min [V]")),
                "v_mpp_max": _to_float(row.get("V MPPT max [V]")),
                "num_mppt": num_mppt,
                "num_inputs": num_inputs,
                "mppt_currents": mppt_currents,
                "lin_max": lin_max,
                "tensao": tensao,
                "fase": str(fase) if fase else None,
                "protection": str(protection) if protection else None,
            }
        )
    return inverters


def run_import(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb_values = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["BD"]
    ws_values = wb_values["BD"]

    modules = parse_modules(ws, ws_values)
    inverters = parse_inverters(ws, ws_values)

    print(f"Modulos encontrados: {len(modules)}")
    print(f"Inversores encontrados: {len(inverters)}")

    Base.metadata.create_all(engine)
    db = SessionLocal()
    try:
        db.query(ModuleRow).delete()
        db.query(InverterRow).delete()
        db.bulk_insert_mappings(ModuleRow, modules)
        db.bulk_insert_mappings(InverterRow, inverters)
        db.commit()
    finally:
        db.close()

    print("Importacao concluida.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python -m app.infra.import_excel <caminho.xlsx>")
        sys.exit(1)
    run_import(sys.argv[1])
